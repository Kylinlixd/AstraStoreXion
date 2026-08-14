#!/usr/bin/env python3
"""Generate deterministic upload fixtures for the blog storage integration.

Run this script with the bundled Codex workspace Python runtime so the DOCX,
PDF, and PNG outputs are reproducible and do not require project dependencies.
"""

from __future__ import annotations

import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "test" / "fixtures" / "generated"
PNG_PATH = OUTPUT_DIR / "storage-test-cover.png"
DOCX_PATH = OUTPUT_DIR / "storage-integration-checklist.docx"
PDF_PATH = OUTPUT_DIR / "storage-client-guide.pdf"
TEXT_PATH = OUTPUT_DIR / "upload-smoke.txt"
MANIFEST_PATH = OUTPUT_DIR / "manifest.json"

FONT_ARIAL_UNICODE = Path("/Library/Fonts/Arial Unicode.ttf")
FONT_HEITI_MEDIUM = Path("/System/Library/Fonts/STHeiti Medium.ttc")
FONT_HEITI_LIGHT = Path("/System/Library/Fonts/STHeiti Light.ttc")

NAVY = "0B1530"
BLUE = "315CF4"
ICE = "EAF0FF"
MUTED = "64748B"
INK = "172033"
PALE = "F5F7FB"
GREEN = "16A085"


def ensure_output_dir() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def load_image_font(size: int, light: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_HEITI_LIGHT if light else FONT_HEITI_MEDIUM
    return ImageFont.truetype(str(path), size=size)


def generate_png() -> None:
    width, height = 1600, 900
    base = Image.new("RGB", (width, height), "#071129")
    pixels = base.load()
    for y in range(height):
        for x in range(width):
            distance = math.hypot(x - width * 0.55, y - height * 0.45)
            glow = max(0.0, 1.0 - distance / 950.0)
            pixels[x, y] = (
                int(7 + 9 * glow),
                int(17 + 20 * glow),
                int(41 + 46 * glow),
            )

    glow_layer = Image.new("RGBA", base.size, (0, 0, 0, 0))
    glow_draw = ImageDraw.Draw(glow_layer)
    for radius, alpha in [(300, 20), (220, 35), (140, 55)]:
        glow_draw.ellipse(
            (1070 - radius, 420 - radius, 1070 + radius, 420 + radius),
            fill=(49, 92, 244, alpha),
        )
    glow_layer = glow_layer.filter(ImageFilter.GaussianBlur(65))
    base = Image.alpha_composite(base.convert("RGBA"), glow_layer)

    draw = ImageDraw.Draw(base)
    for index in range(7):
        x = 930 + index * 76
        draw.line((x, 155, x + 120, 745), fill=(96, 165, 250, 30), width=2)
    for index in range(5):
        y = 250 + index * 95
        draw.line((820, y, 1480, y - 34), fill=(96, 165, 250, 25), width=2)

    # The central crystal and pedestal provide sharp upload/download QA edges.
    crystal = [(1070, 205), (1240, 340), (1160, 585), (970, 585), (900, 342)]
    draw.polygon(crystal, fill="#315CF4")
    draw.polygon([(1070, 205), (1240, 340), (1070, 410)], fill="#6D8BFF")
    draw.polygon([(1070, 205), (1070, 410), (900, 342)], fill="#2846C7")
    draw.polygon([(900, 342), (1070, 410), (970, 585)], fill="#18339E")
    draw.polygon([(1240, 340), (1160, 585), (1070, 410)], fill="#4172FF")
    draw.line(crystal + [crystal[0]], fill="#B7C7FF", width=4)

    draw.rounded_rectangle((880, 630, 1265, 720), radius=18, fill="#101D3D", outline="#4868C6", width=3)
    draw.rounded_rectangle((920, 675, 1225, 750), radius=16, fill="#0C1732", outline="#233A72", width=2)
    for index in range(4):
        draw.ellipse((955 + index * 45, 704, 971 + index * 45, 720), fill="#31D4B2")

    # Orbital paths and file particles are intentionally simple and crisp.
    draw.arc((830, 260, 1320, 620), 195, 515, fill="#91A8FF", width=4)
    draw.arc((865, 285, 1280, 600), 10, 340, fill="#3BDCC1", width=3)
    for x, y in [(852, 432), (1286, 455), (1188, 270), (958, 575)]:
        draw.rounded_rectangle((x, y, x + 34, y + 42), radius=5, fill="#F4F7FF")
        draw.line((x + 9, y + 14, x + 25, y + 14), fill="#315CF4", width=3)
        draw.line((x + 9, y + 23, x + 25, y + 23), fill="#91A8FF", width=3)

    kicker_font = load_image_font(28, light=True)
    title_font = load_image_font(76)
    body_font = load_image_font(30, light=True)
    badge_font = load_image_font(24)
    draw.text((110, 120), "BLOG STORAGE INTEGRATION", font=kicker_font, fill="#6F90FF")
    draw.text((106, 195), "星辰离子 X", font=title_font, fill="#F7FAFF")
    draw.text((110, 305), "上传 · 下载 · 校验", font=body_font, fill="#A8B7D4")

    badges = [("PNG", 110), ("1600 x 900", 250), ("SHA-256", 492)]
    for text, x in badges:
        bbox = draw.textbbox((0, 0), text, font=badge_font)
        box_width = bbox[2] - bbox[0] + 44
        draw.rounded_rectangle((x, 690, x + box_width, 748), radius=16, fill="#122653", outline="#3559B8", width=2)
        draw.text((x + 22, 705), text, font=badge_font, fill="#DCE6FF")

    draw.text((110, 798), "Deterministic visual fixture / v1", font=kicker_font, fill="#71809B")
    base.convert("RGB").save(PNG_PATH, format="PNG", optimize=True)


def set_run_font(run, *, name: str = "Arial Unicode MS", east_asia: str = "Arial Unicode MS", size: float | None = None,
                 color: str | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 140, bottom: int = 100, end: int = 140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_docx_paragraph(doc: Document, text: str, *, size: float = 11, color: str = INK,
                       bold: bool = False, before: float = 0, after: float = 6,
                       align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(text), size=size, color=color, bold=bold)
    return paragraph


def add_docx_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}[level],
            color=BLUE if level < 3 else "1F4D78",
            bold=True,
        )
    paragraph.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[level])
    paragraph.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}[level])
    return paragraph


def add_checklist_table(doc: Document, items: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, text in enumerate(("Done", "Check", "Acceptance")):
        set_cell_shading(header[index], ICE)
        paragraph = header[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(paragraph.add_run(text), size=10, color=NAVY, bold=True)
    for item, standard in items:
        cells = table.add_row().cells
        status = cells[0].paragraphs[0]
        status.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(status.add_run("☐"), size=14, color=BLUE)
        set_run_font(cells[1].paragraphs[0].add_run(item), size=10.5, color=INK, bold=True)
        set_run_font(cells[2].paragraphs[0].add_run(standard), size=10.5, color=INK)
    set_table_geometry(table, [900, 3150, 5310])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def generate_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("ASTRASTORE XION / BLOG STORAGE"), size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("Deployment acceptance  |  "), size=8.5, color=MUTED)
    add_page_field(footer)

    add_docx_paragraph(doc, "OPERATOR CHECKLIST", size=10, color=BLUE, bold=True, after=4)
    add_docx_paragraph(doc, "Blog Storage Deployment Checklist", size=27, color=NAVY, bold=True, after=7)
    add_docx_paragraph(
        doc,
        "Preflight, release verification, and rollback record for AstraStoreXion, the Django blog API, and the Vue admin client.",
        size=12.5,
        color=MUTED,
        after=18,
    )

    metadata = doc.add_table(rows=2, cols=4)
    metadata.style = "Table Grid"
    values = [
        ("Environment", "Production server"),
        ("Release date", "2026-08-10"),
        ("Storage", "New: Xion / Legacy: local media"),
        ("File limit", "1 GB"),
    ]
    for row_index in range(2):
        for pair_index in range(2):
            label, value = values[row_index * 2 + pair_index]
            label_cell = metadata.rows[row_index].cells[pair_index * 2]
            value_cell = metadata.rows[row_index].cells[pair_index * 2 + 1]
            set_cell_shading(label_cell, ICE)
            set_run_font(label_cell.paragraphs[0].add_run(label), size=9.5, color=NAVY, bold=True)
            set_run_font(value_cell.paragraphs[0].add_run(value), size=9.5, color=INK)
    set_table_geometry(metadata, [1500, 1900, 1500, 4460])

    add_docx_heading(doc, "1. Preflight", 1)
    add_checklist_table(
        doc,
        [
            ("Back up code and database", "Backups are timestamped and restorable; the media directory remains untouched."),
            ("Verify service-token permissions", "The token exists only in restricted env files, never in Git, logs, or the browser."),
            ("Confirm disk headroom", "Free space covers current data plus at least one complete release package."),
            ("Run the full verification suite", "Go, Python, Django, and Vue tests plus the production build all pass."),
        ],
    )

    add_docx_heading(doc, "2. Service release", 1)
    add_checklist_table(
        doc,
        [
            ("Start the Xion systemd service", "It listens only on 127.0.0.1; /healthz and /readyz succeed."),
            ("Apply Django migrations", "UploadFile has the new fields; historical rows default to local."),
            ("Enable XION_STORAGE_ENABLED", "The blog reads the service URL, token, and 1 GB limit."),
            ("Switch the frontend release", "The current symlink changes atomically and the prior release remains available."),
        ],
    )

    add_docx_heading(doc, "3. End-to-end acceptance", 1)
    add_checklist_table(
        doc,
        [
            ("Upload PNG", "The UI shows progress, preview works, and the backend reports xion."),
            ("Upload PDF and DOCX", "Both are classified as document; downloads preserve original names."),
            ("Upload TXT", "Size, MIME, and SHA-256 match the local fixture."),
            ("Delete test files", "Database rows and Xion objects are removed; repeated delete is safe."),
            ("Verify legacy media", "Existing /media/ links still work without a bulk migration."),
        ],
    )

    add_docx_heading(doc, "4. Rollback record", 1)
    add_docx_paragraph(
        doc,
        "Trigger rollback when health checks fail, uploads keep returning 5xx, migration fails, or legacy media becomes unavailable.",
        size=10.5,
        color="7A5A00",
        bold=True,
        after=8,
    )
    rollback = doc.add_table(rows=3, cols=2)
    rollback.style = "Table Grid"
    rollback_rows = [
        ("Frontend", "Point current to the previous release, then reload Nginx"),
        ("Backend", "Restore the code and .env backups; roll back the database only if required"),
        ("Storage", "Disable XION_STORAGE_ENABLED; keep the data directory and its objects"),
    ]
    for index, (label, action) in enumerate(rollback_rows):
        set_cell_shading(rollback.rows[index].cells[0], PALE)
        set_run_font(rollback.rows[index].cells[0].paragraphs[0].add_run(label), size=10, color=NAVY, bold=True)
        set_run_font(rollback.rows[index].cells[1].paragraphs[0].add_run(action), size=10, color=INK)
    set_table_geometry(rollback, [1800, 7560])

    doc.core_properties.title = "Blog Storage Deployment Checklist"
    doc.core_properties.subject = "AstraStoreXion integration test fixture"
    doc.core_properties.author = "AstraStoreXion"
    doc.core_properties.keywords = "storage, blog, deployment, test fixture"
    fixed_time = datetime(2026, 8, 10, tzinfo=timezone.utc)
    doc.core_properties.created = fixed_time
    doc.core_properties.modified = fixed_time
    doc.save(DOCX_PATH)


def register_pdf_fonts() -> None:
    pdfmetrics.registerFont(TTFont("ArialUnicode", str(FONT_ARIAL_UNICODE)))
    pdfmetrics.registerFont(TTFont("Heiti", str(FONT_HEITI_MEDIUM)))


def generate_pdf() -> None:
    register_pdf_fonts()
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.78 * inch,
        bottomMargin=0.72 * inch,
        title="AstraStoreXion 博客文件中心使用手册",
        author="AstraStoreXion",
        invariant=1,
    )

    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "GuideTitle",
        parent=styles["Title"],
        fontName="Heiti",
        fontSize=27,
        leading=36,
        textColor=colors.HexColor("#0B1530"),
        alignment=TA_LEFT,
        spaceAfter=12,
    )
    subtitle = ParagraphStyle(
        "GuideSubtitle",
        parent=styles["Normal"],
        fontName="ArialUnicode",
        fontSize=11,
        leading=18,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=18,
    )
    h1 = ParagraphStyle(
        "GuideH1",
        parent=styles["Heading1"],
        fontName="Heiti",
        fontSize=17,
        leading=23,
        textColor=colors.HexColor("#315CF4"),
        spaceBefore=8,
        spaceAfter=10,
        keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "GuideH2",
        parent=styles["Heading2"],
        fontName="Heiti",
        fontSize=12,
        leading=17,
        textColor=colors.HexColor("#0B1530"),
        spaceBefore=8,
        spaceAfter=6,
        keepWithNext=True,
    )
    body = ParagraphStyle(
        "GuideBody",
        parent=styles["BodyText"],
        fontName="ArialUnicode",
        fontSize=10.3,
        leading=17,
        textColor=colors.HexColor("#172033"),
        spaceAfter=8,
    )
    small = ParagraphStyle(
        "GuideSmall",
        parent=body,
        fontSize=8.8,
        leading=14,
        textColor=colors.HexColor("#64748B"),
    )
    callout = ParagraphStyle(
        "GuideCallout",
        parent=body,
        fontSize=10,
        leading=16,
        leftIndent=10,
        rightIndent=10,
        borderColor=colors.HexColor("#B8C7FF"),
        borderWidth=1,
        borderPadding=10,
        backColor=colors.HexColor("#F2F5FF"),
        spaceBefore=4,
        spaceAfter=12,
    )

    def on_page(canvas, pdf_doc) -> None:
        canvas.saveState()
        canvas.setFont("ArialUnicode", 8.5)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(0.75 * inch, 0.42 * inch, "ASTRASTORE XION / BLOG FILE CENTER")
        canvas.drawRightString(7.75 * inch, 0.42 * inch, f"第 {pdf_doc.page} 页")
        canvas.setStrokeColor(colors.HexColor("#DCE3F0"))
        canvas.line(0.75 * inch, 0.58 * inch, 7.75 * inch, 0.58 * inch)
        canvas.restoreState()

    story = [
        Paragraph("CLIENT GUIDE / v1.0", small),
        Spacer(1, 0.12 * inch),
        Paragraph("博客文件中心使用手册", title),
        Paragraph("上传、复制链接、插入文章、下载与安全删除。适用于新 Xion 文件与历史 /media/ 文件并行运行的博客。", subtitle),
        Table(
            [
                [Paragraph("适用对象", h2), Paragraph("博客管理员、内容编辑者、部署运维人员", body)],
                [Paragraph("上传限制", h2), Paragraph("单个文件不超过 1 GB", body)],
                [Paragraph("存储模式", h2), Paragraph("新上传写入 AstraStoreXion，历史本地文件继续可用", body)],
            ],
            colWidths=[1.35 * inch, 5.65 * inch],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0FF")),
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#CFD8EA")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 0.22 * inch),
        Paragraph("快速开始", h1),
        Paragraph("进入管理后台的“文件管理”页面。页面顶部会显示当前可见文件数量、容量和存储兼容状态。", body),
        ListFlowable(
            [
                ListItem(Paragraph("拖拽文件到上传区，或点击“上传文件”。", body)),
                ListItem(Paragraph("等待进度达到 100%，确认列表出现文件名。", body)),
                ListItem(Paragraph("复制公开链接，并粘贴到文章编辑器。", body)),
                ListItem(Paragraph("发布前预览文章，确认图片或下载链接可用。", body)),
            ],
            bulletType="1",
            start="1",
            leftIndent=22,
            bulletFontName="ArialUnicode",
            bulletFontSize=10,
        ),
        Paragraph("重要：不要在文章中使用 127.0.0.1、8081 或服务器内部存储地址。文件链接必须由博客 API 统一提供。", callout),
        PageBreak(),
        Paragraph("支持的文件与常用操作", h1),
        Paragraph("文件中心会根据扩展名和 MIME 自动归类。预览能力取决于浏览器，不能预览不代表文件损坏。", body),
        Table(
            [
                [Paragraph("类型", h2), Paragraph("示例", h2), Paragraph("推荐用途", h2)],
                [Paragraph("图片", body), Paragraph("PNG, JPG, GIF, WebP", body), Paragraph("文章配图、封面、截图", body)],
                [Paragraph("文档", body), Paragraph("PDF, DOCX, XLSX, TXT", body), Paragraph("教程、附件、下载资料", body)],
                [Paragraph("音视频", body), Paragraph("MP3, WAV, MP4, WebM", body), Paragraph("播客、演示、媒体内容", body)],
                [Paragraph("其他", body), Paragraph("ZIP 等", body), Paragraph("仅在业务确有需要时上传", body)],
            ],
            colWidths=[1.15 * inch, 2.15 * inch, 3.7 * inch],
            repeatRows=1,
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF0FF")),
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#CFD8EA")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 9),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 0.2 * inch),
        Paragraph("复制链接与插入文章", h2),
        Paragraph("上传完成后，在文件行选择“复制链接”。图片应放入图片组件或 Markdown 图片语法；PDF、Word 和压缩包使用普通下载链接。", body),
        Paragraph("下载", h2),
        Paragraph("点击“下载”会保留原始文件名。若浏览器直接打开 PDF，可使用浏览器的保存功能。", body),
        Paragraph("删除", h2),
        Paragraph("删除前先搜索博客正文，确认没有文章仍引用该链接。存储服务临时不可用时，系统会保留数据库记录，避免出现无法追踪的残留对象。", body),
        Paragraph("历史文件仍保留在 /media/，无需为了启用 Xion 而批量迁移。", callout),
        PageBreak(),
        Paragraph("故障排查与验收", h1),
        Table(
            [
                [Paragraph("现象", h2), Paragraph("先检查", h2), Paragraph("处理", h2)],
                [Paragraph("上传前即失败", body), Paragraph("是否超过 1 GB，扩展名是否正常", body), Paragraph("压缩或拆分文件，刷新后重试", body)],
                [Paragraph("进度中断", body), Paragraph("网络连接、登录状态、服务健康", body), Paragraph("重新登录；运维检查 Xion 与博客服务", body)],
                [Paragraph("可下载但不能预览", body), Paragraph("浏览器是否支持该格式", body), Paragraph("使用本地应用打开下载文件", body)],
                [Paragraph("历史链接异常", body), Paragraph("/media/ Nginx alias 与文件权限", body), Paragraph("不要迁移或删除旧媒体，恢复 alias 配置", body)],
            ],
            colWidths=[1.45 * inch, 2.55 * inch, 3 * inch],
            repeatRows=1,
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF0FF")),
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#CFD8EA")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 9),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 0.2 * inch),
        Paragraph("上线验收最小集", h2),
        ListFlowable(
            [
                ListItem(Paragraph("PNG、PDF、DOCX、TXT 各上传一次。", body)),
                ListItem(Paragraph("逐个下载并比对 SHA-256。", body)),
                ListItem(Paragraph("确认数据库记录的 storage_backend 为 xion。", body)),
                ListItem(Paragraph("删除测试文件，再次下载应返回 404。", body)),
                ListItem(Paragraph("抽查至少一个历史 /media/ 链接仍返回 200。", body)),
            ],
            bulletType="bullet",
            bulletChar="●",
            leftIndent=22,
            bulletFontName="ArialUnicode",
            bulletFontSize=7,
            bulletColor=colors.HexColor("#315CF4"),
        ),
        Spacer(1, 0.12 * inch),
        KeepTogether(
            [
                Paragraph("安全边界", h2),
                Paragraph("服务密钥不得出现在 Git、浏览器、截图或日志中。Xion 仅监听回环地址，由 Django 负责用户认证和业务元数据。", callout),
            ]
        ),
        Paragraph("文档版本：1.0 / 生成日期：2026-08-10 / 测试资料可在验收完成后删除。", small),
    ]
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def generate_manifest() -> None:
    mime_types = {
        PNG_PATH.name: "image/png",
        PDF_PATH.name: "application/pdf",
        DOCX_PATH.name: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        TEXT_PATH.name: "text/plain",
    }
    entries = []
    for path in (PNG_PATH, PDF_PATH, DOCX_PATH, TEXT_PATH):
        body = path.read_bytes()
        entries.append(
            {
                "name": path.name,
                "mime_type": mime_types[path.name],
                "size": len(body),
                "sha256": hashlib.sha256(body).hexdigest(),
            }
        )
    MANIFEST_PATH.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": "2026-08-10T00:00:00+08:00",
                "files": entries,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


def main() -> None:
    ensure_output_dir()
    generate_png()
    generate_docx()
    generate_pdf()
    generate_manifest()
    for path in (PNG_PATH, DOCX_PATH, PDF_PATH, MANIFEST_PATH):
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
